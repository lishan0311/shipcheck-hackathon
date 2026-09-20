from pathlib import Path
from unittest.mock import patch
import pytest
from backend.documents import read_document, tesseract_command
from backend.comparison import compare_email


def test_docx_preserves_table_and_paragraph_locations(tmp_path):
    from docx import Document
    doc=Document()
    doc.add_paragraph('SHIPPING INSTRUCTION')
    cells=doc.add_table(rows=1,cols=2).rows[0].cells
    cells[0].text='Container Count'
    cells[1].text='3'
    file=tmp_path/'document.docx'; doc.save(file)
    parsed=read_document(file)
    assert 'Container Count: 3' in parsed['text']
    assert parsed['line_sources'][1]['table_row']==1


def test_xlsx_retains_sheet_cells_and_flags_formulas(tmp_path):
    from openpyxl import Workbook
    workbook=Workbook(); sheet=workbook.active
    sheet.append(['SHIPPING INSTRUCTION'])
    sheet.append(['Gross Weight (KG)',22000])
    sheet.append(['Container Count','=1+2'])
    file=tmp_path/'document.xlsx';workbook.save(file)
    parsed=read_document(file)
    assert 'Gross Weight (KG): 22000' in parsed['text']
    assert parsed['line_sources'][1]['cells']==['A2','B2']
    assert parsed['warnings']


def test_actual_text_pdf_reads_with_page_positions():
    files=list(Path('sdoc-hackathon-bundle/attachments').glob('*.pdf'))
    for file in files:
        # Read-only discovery; never use reference labels.
        import pdfplumber
        with pdfplumber.open(file) as pdf:
            has_text=bool(pdf.pages[0].extract_text())
        if has_text:
            parsed=read_document(file)
            assert any(source.get('page')==1 and 'bbox' in source for source in parsed['line_sources'])
            return
    pytest.skip('No text PDF in this checkout')


def test_scanned_pdf_calls_ocr_and_requires_review(tmp_path):
    from PIL import Image
    file=tmp_path/'scan.pdf'
    Image.new('RGB',(300,300),'white').save(file,'PDF')
    with patch('backend.documents.ocr_page',return_value=(['SHIPPING INSTRUCTION','Container Count: 3'],[{'page':1,'method':'tesseract'},{'page':1,'method':'tesseract'}])) as ocr:
        parsed=read_document(file)
    ocr.assert_called_once_with(file,1)
    assert parsed['warnings']
    report=compare_email({'email':{'email_id':'demo'},'documents':[{'path':'scan.pdf','status':'READ',**parsed}]})
    assert report['comparison_status']=='NEEDS_REVIEW'


def test_configured_tesseract_executable_is_used(tmp_path, monkeypatch):
    executable=tmp_path/'tesseract.exe'
    executable.write_bytes(b'placeholder')
    monkeypatch.setenv('TESSERACT_CMD',str(executable))
    assert tesseract_command()==str(executable)


def test_corrupt_document_does_not_crash_ingestion(tmp_path):
    import json
    from backend.ingestion import load_email
    (tmp_path/'inbox').mkdir();(tmp_path/'attachments').mkdir()
    (tmp_path/'attachments/bad.pdf').write_bytes(b'not a PDF')
    (tmp_path/'inbox/demo.json').write_text(json.dumps({'email_id':'demo','from':'x','subject':'x','body':'x','attachments':['attachments/bad.pdf']}))
    document=load_email(tmp_path,'demo')['documents'][0]
    assert document['status']=='UNREADABLE'
    assert 'PDF structure is damaged or incomplete' in document['error']
