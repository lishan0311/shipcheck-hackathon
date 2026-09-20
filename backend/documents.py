"""Document adapters with original page, paragraph and cell locations."""
import os
import shutil
from pathlib import Path


class UnsupportedDocument(ValueError):
    pass


def tesseract_command() -> str:
    """Locate the native OCR executable across local and deployed environments."""
    configured = os.getenv('TESSERACT_CMD', '').strip()
    candidates = [
        configured,
        shutil.which('tesseract') or '',
        r'C:\Program Files\Tesseract-OCR\tesseract.exe',
        r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise RuntimeError(
        'Scanned PDF requires OCR, but the Tesseract executable was not found. '
        'Install Tesseract or set TESSERACT_CMD to tesseract.exe.'
    )


def read_document(path: Path) -> dict:
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError('Attachment exceeds 20 MB.')
    lines, sources, warnings = [], [], []

    def add(text, locator):
        for line in str(text).splitlines():
            lines.append(line)
            sources.append(locator.copy())

    def row_text(values):
        from .comparison import LABELS, clean
        values = [str(value) for value in values if value is not None and str(value).strip()]
        if not values:
            return ''
        first = values[0].rstrip(':').strip()
        if clean(first) in LABELS and len(values) > 1:
            return first + ': ' + ' '.join(values[1:])
        return ' '.join(values)

    suffix = path.suffix.lower()
    if suffix == '.txt':
        for number, line in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1):
            lines.append(line)
            sources.append({'line': number})
    elif suffix == '.docx':
        from docx import Document
        from docx.table import Table
        doc = Document(path)
        for index, block in enumerate(doc.iter_inner_content(), 1):
            if isinstance(block, Table):
                for row, cells in enumerate(block.rows, 1):
                    add(row_text([cell.text for cell in cells.cells]), {'block': index, 'table_row': row})
            else:
                add(block.text, {'paragraph': index})
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook:
                if sheet.max_row * sheet.max_column > 100_000:
                    raise ValueError('Worksheet exceeds the 100,000-cell limit.')
                for row in sheet.iter_rows():
                    populated = [cell for cell in row if cell.value is not None]
                    if not populated:
                        continue
                    if any(cell.data_type == 'f' for cell in populated):
                        warnings.append('Formula cells require review; formulas were not evaluated.')
                    add(row_text([cell.value for cell in populated]), {'sheet': sheet.title, 'cells': [cell.coordinate for cell in populated]})
        finally:
            workbook.close()
    elif suffix == '.pdf':
        import pdfplumber
        try:
            pdf = pdfplumber.open(path)
        except Exception as exc:
            raise ValueError('PDF structure is damaged or incomplete; request a new copy from the sender.') from exc
        with pdf:
            if len(pdf.pages) > 30:
                raise ValueError('PDF exceeds the 30-page processing limit.')
            for number, page in enumerate(pdf.pages, 1):
                extracted = page.extract_text_lines(return_chars=False)
                if extracted:
                    for line in extracted:
                        add(line['text'], {'page': number, 'bbox': [line['x0'], line['top'], line['x1'], line['bottom']], 'method': 'pdfplumber'})
                else:
                    ocr_lines, ocr_sources = ocr_page(path, number)
                    lines.extend(ocr_lines)
                    sources.extend(ocr_sources)
                    # OCR is helpful evidence, not an automatic approval signal.
                    warnings.append(f'Page {number} used OCR; confirm the extracted values before approval.')
    else:
        raise UnsupportedDocument(f'Unsupported attachment format: {suffix or "unknown"}.')
    text = path.read_text(encoding='utf-8-sig') if suffix == '.txt' else '\n'.join(lines)
    if not text.strip():
        raise ValueError('No readable text was found in the attachment.')
    return {'text': text, 'line_sources': sources, 'warnings': warnings}


def ocr_page(path, number):
    import pypdfium2 as pdfium
    import pytesseract
    from pytesseract import Output
    pytesseract.pytesseract.tesseract_cmd = tesseract_command()
    pdf = pdfium.PdfDocument(str(path))
    try:
        page = pdf[number - 1]
        try:
            # Three rendered pixels per PDF point is a useful accuracy gain for
            # small labels and punctuation in scanned shipping forms.
            scale = 3
            if page.get_width() * page.get_height() * scale * scale * 4 > 80_000_000:
                raise ValueError('PDF page is too large for OCR.')
            bitmap = page.render(scale=scale)
            try:
                image = bitmap.to_pil()
                try:
                    data = pytesseract.image_to_data(
                        image,
                        lang=os.getenv('OCR_LANGUAGE', 'eng'),
                        config='--psm 6',
                        output_type=Output.DICT,
                        timeout=45,
                    )
                finally:
                    image.close()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        pdf.close()
    groups = {}
    for i, text in enumerate(data['text']):
        if not text.strip():
            continue
        key = (data['block_num'][i], data['par_num'][i], data['line_num'][i])
        groups.setdefault(key, []).append(i)
    lines, locations = [], []
    for indices in groups.values():
        lines.append(' '.join(data['text'][i] for i in indices))
        confidences = [float(data['conf'][i]) for i in indices if float(data['conf'][i]) >= 0]
        locations.append({'page': number, 'method': 'tesseract', 'confidence': sum(confidences) / len(confidences),
                          'bbox_pixels': [min(data['left'][i] for i in indices), min(data['top'][i] for i in indices),
                                          max(data['left'][i] + data['width'][i] for i in indices), max(data['top'][i] + data['height'][i] for i in indices)]})
    return lines, locations
